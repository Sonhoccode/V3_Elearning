import json
import os
import re
import uuid
from io import BytesIO
from rest_framework import views, status
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.messages import SystemMessage, HumanMessage
from .vector_store import ingest_custom_text, ingest_pdf, delete_documents_by_source, get_supabase_client
from .config import LLM_MODEL_NAME
from .models import InternalDocument
from apps.Common.permissions import IsAdminRole

class AdminIngestView(views.APIView):
    permission_classes = [IsAuthenticated, IsAdminRole]

    def post(self, request):
        input_type = request.data.get("inputType", "text")
        
        try:
            if input_type == "text":
                text = request.data.get("text")
                source_name = request.data.get("source_name", "Admin Input")
                
                if not text:
                    return Response({"error": "Nội dung văn bản bị trống."}, status=status.HTTP_400_BAD_REQUEST)
                    
                ingest_custom_text(text, source_name)
                return Response({"message": "Tài liệu văn bản đã được Vector hóa và lưu thành công!"}, status=status.HTTP_200_OK)
                
            elif input_type == "file":
                file_obj = request.FILES.get("file")
                
                if not file_obj:
                    return Response({"error": "Vui lòng chọn 1 file PDF hoặc DOCX."}, status=status.HTTP_400_BAD_REQUEST)
                filename_lower = file_obj.name.lower()
                if not (filename_lower.endswith(".pdf") or filename_lower.endswith(".docx")):
                    return Response({"error": "Hệ thống chỉ tạm thời hỗ trợ file định dạng .pdf hoặc .docx"}, status=status.HTTP_400_BAD_REQUEST)
                if filename_lower.endswith(".docx"):
                    try:
                        from docx import Document as DocxDocument  # noqa: F401
                    except ImportError:
                        return Response(
                            {"error": "Thiếu thư viện python-docx. Hãy cài `python-docx` để dùng file .docx."},
                            status=status.HTTP_400_BAD_REQUEST,
                        )
                if filename_lower.endswith(".pdf"):
                    try:
                        from google import genai  # noqa: F401
                    except ImportError:
                        return Response(
                            {"error": "Thiếu thư viện google-genai. Hãy cài `google-genai` để dùng OCR PDF."},
                            status=status.HTTP_400_BAD_REQUEST,
                        )
                    if not os.getenv("GOOGLE_API_KEY"):
                        return Response(
                            {"error": "Thiếu GOOGLE_API_KEY để OCR PDF."},
                            status=status.HTTP_400_BAD_REQUEST,
                        )

                file_bytes = file_obj.read()
                bucket = os.getenv("SUPABASE_STORAGE_BUCKET", "internal-docs")
                storage_path = f"legacy/{uuid.uuid4().hex}_{os.path.basename(file_obj.name)}"
                try:
                    supabase = get_supabase_client()
                    supabase.storage.from_(bucket).upload(
                        storage_path,
                        file_bytes,
                        {"content-type": file_obj.content_type or "application/octet-stream"},
                    )
                except Exception as exc:
                    return Response(
                        {"error": f"Không thể upload file lên storage: {exc}"},
                        status=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    )
                
                if filename_lower.endswith(".docx"):
                    try:
                        from docx import Document as DocxDocument
                    except ImportError:
                        return Response(
                            {"error": "Thiếu thư viện python-docx. Hãy cài `python-docx` để dùng file .docx."},
                            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
                        )
                    text_content = []
                    docx = DocxDocument(BytesIO(file_bytes))
                    for para in docx.paragraphs:
                        if para.text:
                            text_content.append(para.text)
                    ingest_custom_text("\n".join(text_content), source_name=file_obj.name)
                else:
                    # Chạy quá trình Embedding của File PDF
                    ingest_pdf(file_bytes=file_bytes, source_key=storage_path, display_name=file_obj.name)
                
                return Response({"message": f"File '{file_obj.name}' đã tải lên và phân tách (chunks) nạp vào DB thành công!"}, status=status.HTTP_200_OK)
            
            else:
                return Response({"error": "Tham số inputType không hợp lệ."}, status=status.HTTP_400_BAD_REQUEST)
                
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class InternalDocumentListCreateView(views.APIView):
    permission_classes = [IsAuthenticated, IsAdminRole]

    def get(self, request):
        docs = InternalDocument.objects.all().order_by("-created_at")
        items = [
            {
                "id": str(doc.id),
                "title": doc.title,
                "source_type": doc.source_type,
                "source_key": doc.source_key,
                "original_filename": doc.original_filename,
                "stored_path": doc.stored_path,
                "chunk_count": doc.chunk_count,
                "status": doc.status,
                "error_message": doc.error_message,
                "created_at": doc.created_at,
                "updated_at": doc.updated_at,
            }
            for doc in docs
        ]
        return Response({"items": items}, status=status.HTTP_200_OK)

    def post(self, request):
        input_type = request.data.get("inputType", "text")
        title = (request.data.get("title") or request.data.get("source_name") or "").strip()

        if input_type not in {"text", "file"}:
            return Response({"error": "Tham số inputType không hợp lệ."}, status=status.HTTP_400_BAD_REQUEST)

        if input_type == "text":
            text = request.data.get("text")
            if not text:
                return Response({"error": "Nội dung văn bản bị trống."}, status=status.HTTP_400_BAD_REQUEST)
            if not title:
                title = "Admin Input"
        else:
            file_obj = request.FILES.get("file")
            if not file_obj:
                return Response({"error": "Vui lòng chọn 1 file PDF hoặc DOCX."}, status=status.HTTP_400_BAD_REQUEST)
            filename_lower = file_obj.name.lower()
            if not (filename_lower.endswith(".pdf") or filename_lower.endswith(".docx")):
                return Response({"error": "Hệ thống chỉ hỗ trợ file .pdf hoặc .docx"}, status=status.HTTP_400_BAD_REQUEST)
            if not title:
                title = os.path.splitext(file_obj.name)[0]
            if filename_lower.endswith(".docx"):
                try:
                    from docx import Document as DocxDocument  # noqa: F401
                except ImportError:
                    return Response(
                        {"error": "Thiếu thư viện python-docx. Hãy cài `python-docx` để dùng file .docx."},
                        status=status.HTTP_400_BAD_REQUEST,
                    )
            if filename_lower.endswith(".pdf"):
                try:
                    from google import genai  # noqa: F401
                except ImportError:
                    return Response(
                        {"error": "Thiếu thư viện google-genai. Hãy cài `google-genai` để dùng OCR PDF."},
                        status=status.HTTP_400_BAD_REQUEST,
                    )
                if not os.getenv("GOOGLE_API_KEY"):
                    return Response(
                        {"error": "Thiếu GOOGLE_API_KEY để OCR PDF."},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

        doc_id = uuid.uuid4()
        source_key = f"internal:{doc_id}"
        doc = InternalDocument(
            id=doc_id,
            title=title,
            source_type="text" if input_type == "text" else "pdf",
            source_key=source_key,
            status="processing",
            created_by=request.user,
        )
        doc.save()

        try:
            if input_type == "text":
                chunk_count = ingest_custom_text(
                    text=text,
                    source_name=title,
                    source_key=source_key,
                    display_name=title,
                    document_id=str(doc_id),
                )
            else:
                file_bytes = file_obj.read()
                bucket = os.getenv("SUPABASE_STORAGE_BUCKET", "internal-docs")
                storage_path = f"{doc_id}/{os.path.basename(file_obj.name)}"
                supabase = get_supabase_client()
                supabase.storage.from_(bucket).upload(
                    storage_path,
                    file_bytes,
                    {"content-type": file_obj.content_type or "application/octet-stream"},
                )

                doc.original_filename = file_obj.name
                doc.stored_path = f"{bucket}/{storage_path}"
                doc.save(update_fields=["original_filename", "stored_path"])
                filename_lower = file_obj.name.lower()
                if filename_lower.endswith(".docx"):
                    try:
                        from docx import Document as DocxDocument
                    except ImportError:
                        return Response(
                            {"error": "Thiếu thư viện python-docx. Hãy cài `python-docx` để dùng file .docx."},
                            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
                        )
                    text_content = []
                    docx = DocxDocument(BytesIO(file_bytes))
                    for para in docx.paragraphs:
                        if para.text:
                            text_content.append(para.text)
                    chunk_count = ingest_custom_text(
                        text="\n".join(text_content),
                        source_name=title,
                        source_key=source_key,
                        display_name=title,
                        document_id=str(doc_id),
                    )
                    doc.source_type = "docx"
                    doc.save(update_fields=["source_type"])
                else:
                    chunk_count = ingest_pdf(
                        file_bytes=file_bytes,
                        source_key=source_key,
                        display_name=title,
                        document_id=str(doc_id),
                    )

            doc.chunk_count = int(chunk_count or 0)
            doc.status = "ready"
            doc.error_message = None
            doc.save(update_fields=["chunk_count", "status", "error_message", "updated_at"])

            return Response(
                {
                    "id": str(doc.id),
                    "title": doc.title,
                    "source_type": doc.source_type,
                    "source_key": doc.source_key,
                    "original_filename": doc.original_filename,
                    "stored_path": doc.stored_path,
                    "chunk_count": doc.chunk_count,
                    "status": doc.status,
                    "error_message": doc.error_message,
                    "created_at": doc.created_at,
                    "updated_at": doc.updated_at,
                },
                status=status.HTTP_201_CREATED,
            )
        except Exception as exc:
            doc.status = "error"
            doc.error_message = str(exc)
            doc.save(update_fields=["status", "error_message", "updated_at"])
            return Response({"error": str(exc)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class InternalDocumentDetailView(views.APIView):
    permission_classes = [IsAuthenticated, IsAdminRole]

    def delete(self, request, doc_id):
        try:
            doc = InternalDocument.objects.get(id=doc_id)
        except InternalDocument.DoesNotExist:
            return Response({"error": "Tài liệu không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        delete_documents_by_source(doc.source_key)

        if doc.stored_path:
            bucket = None
            key = None
            if "/" in doc.stored_path:
                bucket, key = doc.stored_path.split("/", 1)
            if bucket and key:
                try:
                    supabase = get_supabase_client()
                    supabase.storage.from_(bucket).remove([key])
                except Exception:
                    pass

        if doc.stored_path and os.path.exists(doc.stored_path):
            try:
                os.remove(doc.stored_path)
            except Exception:
                pass

        doc.delete()
        return Response({"success": True}, status=status.HTTP_200_OK)


class QuizGenerateView(views.APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        topic = (request.data.get("topic") or "Lập trình cơ bản").strip()
        level = (request.data.get("level") or "").strip()
        num_questions = int(request.data.get("num_questions") or 10)
        choices_per_question = int(request.data.get("choices_per_question") or 4)

        num_questions = max(1, min(num_questions, 20))
        choices_per_question = max(2, min(choices_per_question, 6))

        user_test_result = getattr(request.user, "test_result", None)
        if not level and user_test_result:
            level = user_test_result.level
        if not level:
            level = "Beginner"

        system_prompt = f"""
Bạn là hệ thống tạo đề kiểm tra năng lực.
Hãy tạo {num_questions} câu trắc nghiệm về chủ đề: "{topic}" phù hợp với trình độ: "{level}".
Mỗi câu có đúng {choices_per_question} đáp án.

Trả về DUY NHẤT 1 JSON hợp lệ theo mẫu:
{{
  "questions": [
    {{
      "id": "q1",
      "title": "Câu hỏi ...",
      "answers": [
        {{ "id": "A", "text": "Đáp án A" }},
        {{ "id": "B", "text": "Đáp án B" }},
        ...
      ],
      "correctAnswerId": "A"
    }}
  ]
}}
Không được thêm giải thích hay ký tự thừa ngoài JSON.
"""

        try:
            llm = ChatGoogleGenerativeAI(
                model=LLM_MODEL_NAME,
                temperature=0.2,
                api_key=os.getenv("GOOGLE_API_KEY"),
            )

            response = llm.invoke(
                [SystemMessage(content=system_prompt), HumanMessage(content="Tạo bộ câu hỏi.")]
            )

            raw = response.content if hasattr(response, "content") else str(response)
            json_text = raw
            try:
                data = json.loads(json_text)
            except json.JSONDecodeError:
                match = re.search(r"\{.*\}", raw, re.S)
                if not match:
                    return Response(
                        {"error": "AI trả về dữ liệu không hợp lệ."},
                        status=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    )
                json_text = match.group(0)
                data = json.loads(json_text)

            questions = data.get("questions", [])
            if not isinstance(questions, list) or not questions:
                return Response(
                    {"error": "AI trả về danh sách câu hỏi rỗng."},
                    status=status.HTTP_500_INTERNAL_SERVER_ERROR,
                )

            normalized = []
            for idx, q in enumerate(questions, start=1):
                answers = q.get("answers") or []
                if not isinstance(answers, list) or not answers:
                    continue
                q_id = q.get("id") or f"q{idx}"
                q_title = (q.get("title") or "").strip()
                if not q_title:
                    q_title = f"Câu hỏi {idx}"
                normalized_answers = []
                for a_idx, ans in enumerate(answers):
                    ans_id = ans.get("id") or chr(65 + a_idx)
                    ans_text = (ans.get("text") or "").strip()
                    if not ans_text:
                        ans_text = f"Đáp án {ans_id}"
                    normalized_answers.append({"id": ans_id, "text": ans_text})
                correct_id = q.get("correctAnswerId") or normalized_answers[0]["id"]
                normalized.append(
                    {
                        "id": q_id,
                        "title": q_title,
                        "answers": normalized_answers[:choices_per_question],
                        "correctAnswerId": correct_id,
                    }
                )

            return Response({"questions": normalized}, status=status.HTTP_200_OK)
        except Exception as exc:
            return Response(
                {"error": f"Không thể tạo quiz: {exc}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )
